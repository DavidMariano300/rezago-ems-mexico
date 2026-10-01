"""Genera el manuscrito en Word (.docx) editable, con formato APA 7ª edición.

Salida: salidas/Articulo_APA7.docx

Lee exactamente el mismo `docs/ARTICULO.md` que produce el PDF, de modo que las
dos versiones no pueden desincronizarse: cualquier corrección se hace una vez en
el Markdown y se regeneran ambas.

Formato aplicado: Times New Roman 12 pt, interlineado doble, márgenes de
2.54 cm, encabezado con el título corto a la izquierda y el número de página a
la derecha, sangría de primera línea de 1.27 cm salvo en resumen, abstract,
notas y referencias, y sangría francesa en la lista de referencias.

Las tablas se construyen con los tres filetes horizontales de la norma —arriba,
bajo el encabezado y al cierre— y sin líneas verticales.

Uso:
    .venv/bin/python src/15_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ / "docs" / "ARTICULO.md"
FIGS = RAIZ / "salidas" / "figuras"
SALIDA = RAIZ / "salidas" / "Articulo_APA7.docx"

ENCABEZADO = "REZAGO EN EDUCACIÓN MEDIA SUPERIOR"
SANGRIA = Cm(1.27)
SIN_SANGRIA = {"resumen", "abstract"}


def campo_pagina(parrafo):
    """Inserta el campo PAGE de Word para que el número se actualice solo."""
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = texto
        run._r.append(e)


def inline(parrafo, texto: str, cursiva_base: bool = False) -> None:
    """Traduce **negrita**, *cursiva* y enlaces a secuencias de Word."""
    texto = re.sub(r"\[([^\]]+)\]\((?:[^)]+)\)", r"\1", texto)
    for trozo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("*") and trozo.endswith("*"):
            parrafo.add_run(trozo[1:-1]).italic = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            r = parrafo.add_run(trozo[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        else:
            parrafo.add_run(trozo).italic = cursiva_base


def parrafo(doc, texto="", *, sangria=True, centrado=False, negrita=False,
            cursiva=False, tam=12, espaciado=2.0, francesa=False, espacio_antes=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = (WD_LINE_SPACING.DOUBLE if espaciado == 2.0
                            else WD_LINE_SPACING.SINGLE)
    pf.space_before = Pt(espacio_antes)
    pf.space_after = Pt(0)
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if francesa:
        pf.left_indent = SANGRIA
        pf.first_line_indent = -SANGRIA
    elif sangria:
        pf.first_line_indent = SANGRIA
    if texto:
        if negrita or cursiva:
            r = p.add_run(texto)
            r.bold, r.italic = negrita, cursiva
            r.font.size = Pt(tam)
        else:
            inline(p, texto)
    for r in p.runs:
        r.font.size = Pt(tam)
        r.font.name = "Times New Roman"
    return p


def bordes_tabla(tabla, n_encabezado=1):
    """Deja solo los tres filetes horizontales que pide la norma."""
    tbl = tabla._tbl
    pr = tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single" if lado in ("top", "bottom") else "none")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        bordes.append(e)
    pr.append(bordes)
    # Filete bajo la fila de encabezado.
    for celda in tabla.rows[n_encabezado - 1].cells:
        tcPr = celda._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        b.append(e)
        tcPr.append(b)


def campos_bloque(lineas):
    meta, filas = {}, []
    for l in lineas:
        m = re.match(r"^(\w+):\s*(.*)$", l.strip())
        if m and not l.strip().startswith("|"):
            meta[m.group(1)] = m.group(2)
        elif l.strip().startswith("|"):
            filas.append(l)
    return meta, filas


def escribe_tabla(doc, num, meta, filas_md):
    filas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in filas_md]
    filas = [f for f in filas if not all(set(c) <= set("-: ") for c in f)]
    if not filas:
        return
    parrafo(doc, f"Tabla {num}", sangria=False, negrita=True, espacio_antes=12)
    parrafo(doc, meta.get("titulo", ""), sangria=False, cursiva=True)

    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j, celda in enumerate(fila):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            p.paragraph_format.space_after = Pt(0)
            if j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            inline(p, celda)
            for r in p.runs:
                r.font.size = Pt(12)
                r.font.name = "Times New Roman"
    bordes_tabla(t)
    if meta.get("nota"):
        p = parrafo(doc, "", sangria=False, espaciado=1.0, espacio_antes=6)
        r = p.add_run("Nota. ")
        r.italic = True
        r.font.size, r.font.name = Pt(10), "Times New Roman"
        inline(p, meta["nota"])
        for r in p.runs[1:]:
            r.font.size, r.font.name = Pt(10), "Times New Roman"


def escribe_figura(doc, num, meta):
    ruta = FIGS / meta["archivo"]
    if not ruta.exists():
        return
    parrafo(doc, f"Figura {num}", sangria=False, negrita=True, espacio_antes=12)
    parrafo(doc, meta.get("titulo", ""), sangria=False, cursiva=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(ruta), width=Cm(16.0))
    if meta.get("nota"):
        p = parrafo(doc, "", sangria=False, espaciado=1.0, espacio_antes=6)
        r = p.add_run("Nota. ")
        r.italic = True
        r.font.size, r.font.name = Pt(10), "Times New Roman"
        inline(p, meta["nota"])
        for r in p.runs[1:]:
            r.font.size, r.font.name = Pt(10), "Times New Roman"


def escribe_portada(doc, meta):
    for _ in range(3):
        parrafo(doc, sangria=False)
    parrafo(doc, meta.get("titulo", ""), sangria=False, centrado=True, negrita=True)
    parrafo(doc, sangria=False)
    for clave in ("autor", "afiliacion", "programa", "asesor", "fecha"):
        if meta.get(clave):
            parrafo(doc, meta[clave], sangria=False, centrado=True)
    for _ in range(6):
        parrafo(doc, sangria=False)
    if meta.get("nota"):
        parrafo(doc, "Nota del autor", sangria=False, centrado=True, cursiva=True)
        parrafo(doc, meta["nota"], sangria=False)
    doc.add_page_break()


def main() -> int:
    doc = Document()
    est = doc.styles["Normal"]
    est.font.name = "Times New Roman"
    est.font.size = Pt(12)
    est.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.54)
        s.left_margin = s.right_margin = Cm(2.54)
        enc = s.header.paragraphs[0]
        enc.text = ENCABEZADO + "\t\t"
        campo_pagina(enc)
        enc.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in enc.runs:
            r.font.name, r.font.size = "Times New Roman", Pt(12)

    lineas = FUENTE.read_text(encoding="utf-8").split("\n")
    i, en_refs, en_bloque = 0, False, False

    while i < len(lineas):
        st = lineas[i].strip()

        if st.startswith(":::"):
            cab = st[3:].strip().split()
            if cab:
                tipo, num = cab[0], (cab[1] if len(cab) > 1 else "")
                cuerpo = []
                i += 1
                while i < len(lineas) and lineas[i].strip() != ":::":
                    cuerpo.append(lineas[i])
                    i += 1
                i += 1
                meta, filas = campos_bloque(cuerpo)
                if tipo == "portada":
                    escribe_portada(doc, meta)
                elif tipo == "tabla":
                    escribe_tabla(doc, num, meta, filas)
                elif tipo == "figura":
                    escribe_figura(doc, num, meta)
                continue
            i += 1
            continue

        if not st or st in ("---", "***"):
            i += 1
            continue

        if st.startswith("#### "):
            parrafo(doc, st[5:], sangria=False, negrita=True, cursiva=True,
                    espacio_antes=12)
            i += 1
            continue
        if st.startswith("### "):
            parrafo(doc, st[4:], sangria=False, negrita=True, espacio_antes=12)
            i += 1
            continue
        if st.startswith("## "):
            titulo = st[3:]
            parrafo(doc, titulo, sangria=False, centrado=True, negrita=True,
                    espacio_antes=12)
            en_refs = titulo.strip().lower().startswith("referencias")
            en_bloque = titulo.strip().lower() in SIN_SANGRIA
            i += 1
            continue
        if st.startswith("# "):
            i += 1
            continue

        # Listas, con las líneas de continuación absorbidas en el mismo elemento.
        marca = re.match(r"^(\d+)\. |^[-*] ", st)
        if marca and not st.startswith("**"):
            k = 0
            while i < len(lineas):
                s = lineas[i].strip()
                m = re.match(r"^(?:(\d+)\.|[-*]) ", s)
                if not s or not m:
                    break
                k += 1
                texto = s[m.end():]
                j = i + 1
                while (j < len(lineas) and lineas[j].strip()
                       and not re.match(r"^(?:\d+\.|[-*]) ", lineas[j].strip())
                       and not lineas[j].strip().startswith((":::", "#", "|"))):
                    texto += " " + lineas[j].strip()
                    j += 1
                vinyeta = f"{k}. " if m.group(1) else "— "
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
                pf.left_indent, pf.first_line_indent = SANGRIA, -SANGRIA
                pf.space_after = Pt(0)
                p.add_run(vinyeta)
                inline(p, texto)
                for r in p.runs:
                    r.font.size, r.font.name = Pt(12), "Times New Roman"
                i = j
            continue

        bloque = []
        while i < len(lineas) and lineas[i].strip():
            s = lineas[i].strip()
            if s.startswith((":::", "#", "|")):
                break
            if re.match(r"^\d+\. |^[-*] ", s) and not s.startswith("**"):
                break
            bloque.append(s)
            i += 1
        if bloque:
            parrafo(doc, " ".join(bloque),
                    sangria=not (en_refs or en_bloque), francesa=en_refs)
        else:
            i += 1

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SALIDA)
    print(f"Escrito: {SALIDA.relative_to(RAIZ)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
